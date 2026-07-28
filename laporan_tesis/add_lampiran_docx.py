# -*- coding: utf-8 -*-
"""Menambahkan bagian LAMPIRAN (tautan Google Drive + rincian isi) ke naskah tesis.

Jalankan setelah Revisi-Tesis-Mustari.docx DITUTUP dari Microsoft Word:
    python add_lampiran_docx.py

Ganti nilai LINK_DRIVE di bawah dengan URL folder Google Drive yang sebenarnya,
lalu jalankan ulang skrip ini (skrip akan menimpa bagian LAMPIRAN yang lama).
"""
import os
import shutil
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

LINK_DRIVE = '[ISI LINK GOOGLE DRIVE DI SINI]'

BASE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(BASE, 'Revisi-Tesis-Mustari.docx')
BACKUP = os.path.join(BASE, 'Revisi-Tesis-Mustari.BACKUP.docx')

ISI_ZIP = [
    ('src/', 'Kode sumber aplikasi SIMTA (komponen Vue, views, router, store Pinia, dan lapisan layanan data asinkron)'),
    ('public/', 'Berkas aset statis aplikasi'),
    ('index.baseline.html, index.optimized.html', 'Titik masuk aplikasi SIMTA versi baseline dan optimized'),
    ('index.cp.baseline.html, index.cp.optimized.html', 'Titik masuk aplikasi Company Profile kedua versi'),
    ('vite.config.baseline.js, vite.config.optimized.js', 'Konfigurasi build Vite untuk SIMTA, berisi implementasi manualChunks sebagaimana dibahas pada Sub-bab 3.1'),
    ('vite.config.cp.baseline.js, vite.config.cp.optimized.js', 'Konfigurasi build Vite untuk Company Profile'),
    ('dist-baseline/, dist-optimized/', 'Hasil kompilasi SIMTA kedua versi, sebagai bukti perbandingan ukuran bundle'),
    ('dist-cp-baseline/, dist-cp-optimized/', 'Hasil kompilasi Company Profile kedua versi'),
    ('ukur_performa.cjs, ukur_performa_cp.cjs', 'Skrip Puppeteer untuk pengukuran PerformanceObserver (uji tunggal)'),
    ('ukur_multirun.cjs, ukur_multirun_cp.cjs', 'Skrip Puppeteer untuk pengukuran 5 repetisi per skenario'),
    ('ukur_lighthouse.cjs', 'Skrip audit otomatis Google Lighthouse v11.x'),
    ('capture_screenshots.cjs', 'Skrip pengambilan tangkapan layar untuk verifikasi kesamaan tampilan (Gambar 3.2 dan 3.3)'),
    ('analisis_statistik.py', 'Skrip perhitungan rata-rata, standar deviasi, dan persentase perbaikan'),
    ('create_final_charts.py, process_visuals.py', 'Skrip pembangkit grafik perbandingan (Gambar 3.4 sampai 3.9)'),
    ('data_pengukuran/', 'Data mentah hasil pengukuran dalam format JSON (21 berkas)'),
    ('package.json, package-lock.json', 'Daftar dependensi beserta versi terkunci, untuk menjamin reprodusibilitas'),
]

ISI_DATA = [
    ('multirun_baseline_ideal.json, multirun_optimized_ideal.json', 'Data 5 repetisi SIMTA pada kondisi normal (Tabel 3.1, 3.2, 3.4)'),
    ('multirun_baseline_cpu_lambat.json, multirun_optimized_cpu_lambat.json', 'Data 5 repetisi SIMTA pada kondisi CPU diperlambat 4x (Tabel 3.3, 3.4)'),
    ('multirun_cp_baseline_ideal.json, multirun_cp_optimized_ideal.json', 'Data 5 repetisi Company Profile pada kondisi normal (Tabel 3.5)'),
    ('multirun_cp_baseline_cpu_lambat.json, multirun_cp_optimized_cpu_lambat.json', 'Data 5 repetisi Company Profile pada kondisi CPU diperlambat 4x (Tabel 3.5)'),
    ('lighthouse_baseline.json, lighthouse_optimized.json', 'Hasil audit Lighthouse SIMTA (Tabel 3.6)'),
    ('lighthouse_cp_baseline.json, lighthouse_cp_optimized.json', 'Hasil audit Lighthouse Company Profile (Tabel 3.7)'),
    ('summary_stats.json', 'Rekapitulasi rata-rata dan standar deviasi seluruh skenario'),
]


def hapus_lampiran_lama(doc):
    """Buang seluruh paragraf/tabel mulai dari heading 'LAMPIRAN' sampai akhir dokumen."""
    body = doc.element.body
    mulai = None
    for i, child in enumerate(body):
        teks = ''.join(child.itertext()).strip().upper()
        if child.tag.endswith('}p') and teks.startswith('LAMPIRAN') and len(teks) < 40:
            mulai = i
            break
    if mulai is None:
        return False
    for child in list(body)[mulai:]:
        if child.tag.endswith('}sectPr'):
            continue
        body.remove(child)
    return True


def _format_tabel(t):
    """Meniru format tabel yang sudah dipakai di naskah: lebar 9345 dxa, rata
    tengah, garis tunggal di seluruh sisi. Naskah ini tidak memiliki style
    bernama 'Table Grid', jadi border dipasang langsung pada tblPr."""
    tblPr = t._tbl.tblPr

    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), '9345')
    tblW.set(qn('w:type'), 'dxa')

    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    tblPr.append(jc)

    borders = OxmlElement('w:tblBorders')
    for sisi in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + sisi)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'auto')
        borders.append(el)
    tblPr.append(borders)


def _arsir(sel, warna='D9D9D9'):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), warna)
    sel._tc.get_or_add_tcPr().append(shd)


def tabel(doc, judul_kolom, baris):
    t = doc.add_table(rows=1, cols=2)
    _format_tabel(t)
    for sel, teks in zip(t.rows[0].cells, judul_kolom):
        sel.text = ''
        run = sel.paragraphs[0].add_run(teks)
        run.bold = True
        _arsir(sel)
    for kiri, kanan in baris:
        sel = t.add_row().cells
        sel[0].text = kiri
        sel[1].text = kanan
    return t


def main():
    if not os.path.exists(BACKUP):
        shutil.copy2(DOCX, BACKUP)
        print('Cadangan dibuat:', os.path.basename(BACKUP))

    doc = Document(DOCX)
    if hapus_lampiran_lama(doc):
        print('Bagian LAMPIRAN lama dihapus, akan ditulis ulang.')

    doc.add_page_break()
    doc.add_heading('LAMPIRAN', level=1)

    doc.add_heading('Lampiran 1. Source Code, Data Pengukuran, dan Bahan Presentasi', level=2)
    doc.add_paragraph(
        'Seluruh kode program aplikasi uji, skrip pengukuran, data mentah hasil pengujian, '
        'serta berkas presentasi tesis ini disimpan secara daring dan dapat diakses melalui tautan berikut:'
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(LINK_DRIVE)
    r.bold = True
    r.font.size = Pt(12)

    doc.add_paragraph(
        'Tautan tersebut mengarah pada satu folder Google Drive yang berisi dua berkas, yaitu '
        'mustari-pnup-main.zip (arsip repositori penelitian: kode program, konfigurasi build, skrip '
        'pengukuran, dan data mentah hasil pengujian) serta Presentasi-Tesis-Mustari.pptx (berkas slide '
        'ujian tesis).'
    )

    doc.add_heading('Rincian Isi Berkas Arsip (ZIP)', level=3)
    doc.add_paragraph('Berkas arsip repositori tersusun dalam struktur berikut:')
    tabel(doc, ('Folder / Berkas', 'Isi'), ISI_ZIP)

    doc.add_heading('Rincian Data Mentah Pengukuran', level=3)
    doc.add_paragraph(
        'Folder data_pengukuran/ memuat data hasil pengujian yang menjadi dasar seluruh tabel dan grafik pada BAB III:'
    )
    tabel(doc, ('Berkas', 'Isi'), ISI_DATA)

    doc.add_heading('Cara Menjalankan Ulang Pengujian', level=3)
    doc.add_paragraph('Untuk mereproduksi hasil pengukuran pada BAB III, langkah-langkahnya adalah sebagai berikut:')
    for baris in [
        'npm install',
        'npx vite build --config vite.config.baseline.js',
        'npx vite build --config vite.config.optimized.js',
        'node ukur_multirun.cjs',
        'node ukur_lighthouse.cjs',
        'python analisis_statistik.py',
        'python create_final_charts.py',
    ]:
        par = doc.add_paragraph()
        run = par.add_run(baris)
        run.font.name = 'Consolas'
        run.font.size = Pt(10)

    doc.add_paragraph('Hasil pengujian akan tersimpan dalam format JSON pada folder data_pengukuran/.')

    doc.save(DOCX)
    print('Selesai. LAMPIRAN ditambahkan ke', os.path.basename(DOCX))
    if LINK_DRIVE.startswith('['):
        print('CATATAN: LINK_DRIVE masih placeholder. Ganti di baris atas skrip lalu jalankan ulang.')


if __name__ == '__main__':
    main()
