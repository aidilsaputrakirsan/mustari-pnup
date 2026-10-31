"""
Menggabungkan naskah revisi penguji (studi ablasi) ke dalam naskah tesis.

Sumber : Revisi-Tesis-Mustari.docx (tidak diubah)
Keluaran: Revisi-Tesis-Mustari-v3-Ablasi.docx

Isi sisipan dibangkitkan oleh buat_revisi_ablasi_docx.py dalam mode cetak
(tanpa kotak "Penempatan"), lalu dipindahkan ke posisinya masing-masing:
  - 2.1.10 setelah 2.1.9, 2.8.1 di akhir 2.8, 3.10 setelah 3.9
  - butir kesimpulan dan saran ke-6, entri Deutsch (1996) di Daftar Pustaka
  - baris baru pada tabel Lampiran 1 dan Lampiran 2, perintah ablasi pada cara menjalankan ulang
"""
import copy
import os

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table

import buat_revisi_ablasi_docx as gen

HERE = os.path.dirname(os.path.abspath(__file__))
SUMBER = os.path.join(HERE, "Revisi-Tesis-Mustari.docx")
KELUARAN = os.path.join(HERE, "Revisi-Tesis-Mustari-v3-Ablasi.docx")


def cari(doc, awalan):
    hasil = [p for p in doc.paragraphs if p.text.strip().startswith(awalan)]
    if len(hasil) != 1:
        raise RuntimeError(f"Penanda '{awalan}' ditemukan {len(hasil)} kali, harus tepat 1")
    return hasil[0]


def sisipkan_setelah(doc, jangkar, bangun):
    """Panggil `bangun(doc)` (menambah elemen di akhir body) lalu pindahkan hasilnya ke setelah `jangkar`."""
    body = doc.element.body
    # Elemen baru selalu disisipkan tepat sebelum sectPr (elemen terakhir body)
    n_sebelum = len(body)
    bangun(doc)
    baru = list(body)[n_sebelum - 1:-1]
    posisi = jangkar
    for el in baru:
        posisi.addnext(el)
        posisi = el
    return posisi


def ganti_teks(p, teks, tebal_awal=None):
    """Isi ulang paragraf hasil salinan dengan teks baru, mempertahankan format run pertama."""
    runs = p.findall(qn("w:r"))
    for r in runs[1:]:
        p.remove(r)
    for r in p.findall(qn("w:hyperlink")):
        p.remove(r)
    from docx.text.paragraph import Paragraph
    para = Paragraph(p, None)
    r0 = para.runs[0]
    r0.text = ""
    if tebal_awal:
        r0.text = tebal_awal
        r0.bold = True
    gen.tambah_teks(para, teks)
    for r in para.runs[1:]:
        r.font.name = r0.font.name
        r.font.size = r0.font.size
    return para


def butir_lanjutan(doc, contoh, teks):
    """Salin paragraf daftar bernomor `contoh` (nomor otomatis berlanjut) dan isi dengan `teks`."""
    el = copy.deepcopy(contoh._p)
    contoh._p.addnext(el)
    ganti_teks(el, teks)
    return el


def tambah_baris(tabel, isi):
    baris = copy.deepcopy(tabel.rows[-1]._tr)
    tabel._tbl.append(baris)
    for sel, teks in zip(tabel.rows[-1].cells, isi):
        p = sel.paragraphs[0]
        for ekstra in sel.paragraphs[1:]:
            ekstra._p.getparent().remove(ekstra._p)
        ganti_teks(p._p, teks)


def tabel_setelah(doc, paragraf):
    el = paragraf._p.getnext()
    while el is not None and el.tag != qn("w:tbl"):
        el = el.getnext()
    return Table(el, paragraf._parent)


def main():
    gen.CETAK = True
    doc = Document(SUMBER)

    # 1.7 Sistematika Penulisan
    p = cari(doc, "BAB III HASIL DAN PEMBAHASAN: Menyajikan")
    p.add_run(" Bab ini juga memuat studi ablasi yang membandingkan kontribusi masing-masing teknik "
              "(code splitting, lazy loading, prefetching, dan Gzip).")

    # BAB II: 2.1.10 sebelum 2.2, 2.8.1 di akhir 2.8
    sisipkan_setelah(doc, cari(doc, "Strategi optimasi yang efektif untuk aplikasi kompleksitas rendah")._p,
                     gen.bagian_teori)
    sisipkan_setelah(doc, cari(doc, "Setiap skenario dijalankan 5 kali")._p, gen.bagian_metode)

    # BAB III: 3.10 setelah paragraf terakhir 3.9
    sisipkan_setelah(doc, cari(doc, "Gambar 3.12 di atas merupakan rekaman")._p, gen.bagian_hasil)

    # BAB IV: butir ke-6 kesimpulan dan saran
    butir_lanjutan(doc, cari(doc, "Teknik ini tidak cocok untuk semua jenis website"), gen.NARASI["kesimpulan"])
    butir_lanjutan(doc, cari(doc, "Penggunaan Machine Learning untuk automated code splitting"), gen.NARASI["saran"])

    # Daftar Pustaka: Deutsch (1996) setelah Bogner & Merkel (urut alfabet)
    bogner = cari(doc, "Bogner, J., & Merkel, M. (2022)")
    el = copy.deepcopy(bogner._p)
    bogner._p.addnext(el)
    ganti_teks(el, "Deutsch, P. (1996). _GZIP File Format Specification Version 4.3_ (RFC 1952). Internet Engineering "
                   "Task Force. https://doi.org/10.17487/RFC1952")

    # Lampiran 1: rincian berkas arsip
    t_arsip = tabel_setelah(doc, cari(doc, "Berkas arsip repositori tersusun dalam struktur berikut"))
    ncol = len(t_arsip.columns)
    for isi in [
        ["`vite.config.ablasi.js`, `src/ablasi/`", "Konfigurasi build dan source enam varian studi ablasi (Sub-bab 2.8.1)"],
        ["`build_ablasi.cjs`, `ukur_ablasi.cjs`", "Skrip build seluruh varian ablasi dan pengukurannya (PerformanceObserver, navigasi, Lighthouse)"],
        ["`analisis_ablasi.py`", "Rekap statistik studi ablasi dan pembangkit Gambar 3.13–3.14"],
    ]:
        tambah_baris(t_arsip, (isi + [""] * ncol)[:ncol])

    # Lampiran 1: rincian data mentah
    t_data = tabel_setelah(doc, cari(doc, "Folder data_pengukuran/ memuat data hasil pengujian"))
    ncol = len(t_data.columns)
    tambah_baris(t_data, (["`ablasi/po_*.json`, `ablasi/lh_*.json`, `ablasi/ringkasan_ablasi.json`",
                           "Data mentah dan ringkasan studi ablasi (Tabel 3.9–3.15)"] + [""] * ncol)[:ncol])

    # Lampiran 1: cara menjalankan ulang
    contoh = cari(doc, "python create_final_charts.py")
    posisi = contoh._p
    for perintah in ["node build_ablasi.cjs", "node ukur_ablasi.cjs", "python analisis_ablasi.py"]:
        el = copy.deepcopy(contoh._p)
        posisi.addnext(el)
        ganti_teks(el, f"`{perintah}`")
        posisi = el

    # Lampiran 2: perubahan isi naskah
    t_ubah = tabel_setelah(doc, cari(doc, "C. Perubahan Isi Naskah"))
    tambah_baris(t_ubah, [
        "Sub-bab 2.1.10, 2.8.1, 3.10; Tabel 2.1, 3.9–3.15; Gambar 3.13–3.14",
        "Belum ada. Keempat teknik hanya diuji bersamaan sehingga kontribusi masing-masing tidak dapat dipisahkan.",
        "Ditambahkan studi ablasi (revisi penguji) yang mengukur kontribusi code splitting, lazy loading, prefetching, "
        "dan Gzip secara terpisah, beserta butir kesimpulan dan saran ke-6.",
    ])

    tambah_baris(t_ubah, [
        "Sub-bab 3.1, cuplikan kode manualChunks",
        "Urutan aturan ditulis memeriksa chart.js lebih dulu, berbeda dari berkas konfigurasi yang dipakai saat pengukuran.",
        "Diselaraskan dengan vite.config.optimized.js yang menghasilkan data Bab III (aturan vendor-vue diperiksa lebih "
        "dulu); dampaknya dibahas pada Sub-bab 3.10.",
    ])

    # Sub-bab 3.1: samakan urutan aturan manualChunks dengan vite.config.optimized.js
    awal = cari(doc, "if (id.includes('node_modules/chart.js') ||")
    baris_kode = [
        "          if (id.includes('node_modules/vue') ||",
        "              id.includes('node_modules/pinia')) {",
        "            return 'vendor-vue'",
        "          }",
        "          if (id.includes('node_modules/chart.js') ||",
        "              id.includes('node_modules/vue-chartjs')) {",
        "            return 'vendor-chart'",
        "          }",
    ]
    el = awal._p
    for teks in baris_kode:
        from docx.text.paragraph import Paragraph
        p = Paragraph(el, None)
        p.runs[0].text = teks
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        el = el.getnext()

    # Siap cetak: hapus seluruh stabilo dari revisi sebelumnya
    for hl in list(doc.element.body.iter(qn("w:highlight"))):
        hl.getparent().remove(hl)

    doc.save(KELUARAN)
    print("Tersimpan:", KELUARAN)


if __name__ == "__main__":
    main()
