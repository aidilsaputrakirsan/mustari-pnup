# -*- coding: utf-8 -*-
"""Menyisipkan ABSTRAK (Indonesia) dan ABSTRACT (Inggris) ke Revisi-Tesis-Mustari.docx,
diletakkan antara halaman judul dan BAB I.

Seluruh angka pada abstrak diambil dari hasil pengukuran akhir yang dilaporkan
pada BAB III (summary_stats.json).

Jalankan saat dokumen tertutup dari Microsoft Word:
    python tambah_abstrak_docx.py
"""
import copy
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.text.paragraph import Paragraph

if sys.stdout.encoding and sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

BASE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(BASE, 'Revisi-Tesis-Mustari.docx')

ABSTRAK_ID = (
    'Single Page Application (SPA) menghadapi permasalahan ukuran bundle JavaScript yang besar '
    'sehingga memperlambat pemuatan awal, khususnya pada perangkat berspesifikasi rendah. '
    'Penelitian ini menguji apakah efektivitas kombinasi hybrid lazy loading dan code splitting '
    'pada SPA berbasis Vue.js dan Vite dipengaruhi oleh tingkat kompleksitas aplikasi. Dua aplikasi '
    'dengan tingkat kompleksitas berbeda dibandingkan, yaitu Sistem Informasi Manajemen Tugas Akhir '
    '(SIMTA) yang memuat pustaka berat Chart.js, Pinia, dan Vue Router, serta sebuah Company Profile '
    'berkonten dominan statis. Masing-masing dikompilasi dalam dua versi, yaitu baseline dengan '
    'eager loading monolitik dan optimized yang menerapkan code splitting melalui manualChunks, '
    'lazy loading berbasis dynamic import, prefetching melalui requestIdleCallback, serta kompresi '
    'Brotli dan Gzip. Pengukuran dilakukan menggunakan W3C PerformanceObserver dan Google Lighthouse '
    'melalui Puppeteer sebanyak lima repetisi pada kondisi normal dan kondisi CPU diperlambat empat '
    'kali. Hasil penelitian menunjukkan bahwa pada SIMTA seluruh metrik membaik, yaitu skor Lighthouse '
    'naik 32,0% dari 56,8 menjadi 75,0, First Contentful Paint turun 34,5%, Largest Contentful Paint '
    'turun 27,5%, dan Total Blocking Time turun 30,6%. Sebaliknya pada Company Profile manfaatnya '
    'tidak signifikan, bahkan First Contentful Paint dan Largest Contentful Paint sedikit memburuk '
    'masing-masing sebesar 15,9% dan 18,0% akibat tambahan permintaan HTTP untuk setiap chunk. '
    'Penelitian ini menyimpulkan bahwa tingkat kompleksitas aplikasi merupakan faktor penentu '
    'efektivitas strategi optimasi, dengan ambang praktis penerapan pada bundle awal yang melebihi '
    '200 KB terkompresi dan memuat pustaka berat yang tidak dibutuhkan pada halaman pertama.'
)

KATA_KUNCI = ('Kata kunci: single page application, code splitting, lazy loading, prefetching, '
              'Vue.js, optimasi performa web')

ABSTRACT_EN = (
    'Single Page Applications (SPA) suffer from large JavaScript bundle sizes that slow down initial '
    'loading, particularly on low-specification devices. This study examines whether the effectiveness '
    'of combined hybrid lazy loading and code splitting in Vue.js and Vite-based SPAs is influenced by '
    'the level of application complexity. Two applications of differing complexity were compared: a '
    'Final Project Management Information System (SIMTA) that relies on heavy libraries including '
    'Chart.js, Pinia, and Vue Router, and a Company Profile dominated by static content. Each was '
    'compiled into two versions: a baseline using monolithic eager loading, and an optimized version '
    'applying code splitting through manualChunks, lazy loading via dynamic imports, prefetching '
    'through requestIdleCallback, and Brotli and Gzip compression. Measurements were conducted using '
    'the W3C PerformanceObserver and Google Lighthouse through Puppeteer across five repetitions under '
    'normal conditions and under fourfold CPU throttling. The results show that all metrics improved '
    'for SIMTA: the Lighthouse performance score increased by 32.0% from 56.8 to 75.0, First '
    'Contentful Paint decreased by 34.5%, Largest Contentful Paint by 27.5%, and Total Blocking Time '
    'by 30.6%. In contrast, the Company Profile gained no significant benefit, with First Contentful '
    'Paint and Largest Contentful Paint slightly deteriorating by 15.9% and 18.0% respectively due to '
    'the additional HTTP requests required for each chunk. This study concludes that application '
    'complexity is a determining factor in the effectiveness of the optimization strategy, with a '
    'practical adoption threshold at initial bundles exceeding 200 KB compressed that contain heavy '
    'libraries not required on the first page.'
)

KEYWORDS = ('Keywords: single page application, code splitting, lazy loading, prefetching, '
            'Vue.js, web performance optimization')


def kuning(par):
    for r in par.runs:
        if r.text.strip():
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def sisip(anchor, gaya_src, teks, align=None, italic=False, bold=False):
    el = copy.deepcopy(gaya_src._p)
    anchor._p.addnext(el)
    par = Paragraph(el, anchor._parent)
    for r in list(par.runs):
        r._element.getparent().remove(r._element)
    run = par.add_run(teks)
    run.italic = italic
    run.bold = bold
    if align is not None:
        par.paragraph_format.alignment = align
    kuning(par)
    return par


def main():
    doc = Document(DOCX)
    P = doc.paragraphs

    if any(p.text.strip().upper() == 'ABSTRAK' for p in P):
        print('ABSTRAK sudah ada. Tidak disisipkan ulang.')
        return

    bab1 = next(p for p in P if p.text.strip() == 'BAB I PENDAHULUAN')
    gaya_h1 = bab1
    gaya_isi = next(p for p in P if p.text.startswith('Single Page Application (SPA) telah menjadi'))

    # sisip tepat sebelum BAB I: pakai paragraf kosong terakhir sebelum bab1
    idx = P.index(bab1)
    anchor = P[idx - 1]

    a = sisip(anchor, gaya_h1, 'ABSTRAK', align=WD_ALIGN_PARAGRAPH.CENTER)
    a.style = gaya_h1.style
    a = sisip(a, gaya_isi, ABSTRAK_ID, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    a = sisip(a, gaya_isi, KATA_KUNCI, align=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=True)
    print('[OK] ABSTRAK (Indonesia) + kata kunci disisipkan')

    b = sisip(a, gaya_h1, 'ABSTRACT', align=WD_ALIGN_PARAGRAPH.CENTER)
    b.style = gaya_h1.style
    b = sisip(b, gaya_isi, ABSTRACT_EN, align=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=True)
    b = sisip(b, gaya_isi, KEYWORDS, align=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=True)
    print('[OK] ABSTRACT (Inggris) + keywords disisipkan')

    # pastikan BAB I mulai di halaman baru
    bab1.paragraph_format.page_break_before = True
    print('[OK] BAB I diatur mulai di halaman baru')

    doc.save(DOCX)
    print('\nTersimpan.')


if __name__ == '__main__':
    main()
