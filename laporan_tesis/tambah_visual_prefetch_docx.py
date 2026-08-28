# -*- coding: utf-8 -*-
"""Menyisipkan Sub-bab 3.9 (verifikasi mekanisme prefetching) beserta tiga gambar
pendukung ke dalam Revisi-Tesis-Mustari.docx.

Gambar dinomori 3.10 sampai 3.12 dan diletakkan setelah Sub-bab 3.8, sehingga
penomoran Gambar 3.1 sampai 3.9 yang sudah ada tidak perlu diubah.

Jalankan saat dokumen tertutup dari Microsoft Word:
    python tambah_visual_prefetch_docx.py
"""
import copy
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Inches

if sys.stdout.encoding and sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

BASE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(BASE, 'Revisi-Tesis-Mustari.docx')
IMG = os.path.join(BASE, 'chapters', 'images')
LEBAR = Inches(6.17)          # samakan dengan gambar yang sudah ada di naskah

JUDUL = '3.9 Verifikasi Mekanisme Prefetching'

INTRO = (
    'Sub-bab 3.2 menjelaskan bahwa prefetching ditambahkan untuk menutup kelemahan lazy loading, '
    'yaitu munculnya jeda pada setiap perpindahan halaman. Bagian ini menyajikan bukti bahwa '
    'mekanisme tersebut benar-benar berjalan pada aplikasi yang diuji, mulai dari perbandingan '
    'konsep, struktur chunk hasil build, hingga rekaman permintaan berkas yang sesungguhnya '
    'terjadi di peramban.'
)

GBR = [
    ('diagram_strategi_pemuatan.png',
     'Gambar 3.10 Perbandingan konseptual tiga strategi pemuatan modul.',
     'Gambar 3.10 di atas membandingkan tiga strategi secara berdampingan. Pada eager loading, '
     'seluruh kode diunduh sekaligus di awal sehingga perpindahan halaman memang terasa instan, '
     'tetapi beban pemuatan awal menjadi besar. Pada lazy loading murni, beban awal berhasil '
     'diperkecil, namun setiap klik menu memunculkan jeda karena chunk halaman baru diunduh pada '
     'saat itu juga. Strategi ketiga, yaitu lazy loading yang dilengkapi prefetching, '
     'mempertahankan bundle awal yang kecil sekaligus menghilangkan jeda tersebut dengan '
     'memindahkan pengunduhan ke waktu senggang peramban.'),
    ('diagram_arsitektur_chunk.png',
     'Gambar 3.11 Peta chunk hasil build SIMTA versi optimized beserta waktu pemuatannya.',
     'Gambar 3.11 di atas memetakan sembilan chunk hasil build ke dalam tiga kelompok berdasarkan '
     'kapan masing-masing diunduh. Kelompok pertama berisi berkas yang dibutuhkan halaman pertama '
     'dengan total 120,1 KB setelah dikompresi. Kelompok kedua berisi dua chunk yang dijadwalkan '
     'melalui prefetching, hanya 5,4 KB terkompresi — biaya yang sangat kecil dibanding beban '
     'pemuatan awal. Kelompok ketiga berisi chunk yang sengaja tidak di-prefetch dan baru diunduh '
     'apabila rutenya benar-benar diakses, sehingga prefetching tetap bersifat selektif dan tidak '
     'mengembalikan aplikasi ke pola eager loading.'),
    ('chart_prefetch_network.png',
     'Gambar 3.12 Rekaman permintaan berkas JavaScript pada versi optimized tanpa interaksi klik '
     '(rata-rata 5 repetisi).',
     'Gambar 3.12 di atas merupakan rekaman permintaan berkas yang sesungguhnya terjadi ketika '
     'aplikasi dibuka lalu dibiarkan tanpa satu pun klik. Terlihat tiga tahap yang berurutan: '
     'berkas entry beserta vendor-vue.js dan vendor-chart.js diminta pada milidetik-milidetik '
     'pertama, DashboardView.js sebagai rute pembuka menyusul pada sekitar 129 ms, kemudian '
     'DaftarJudulView.js dan DetailBimbinganView.js diunduh pada sekitar 351 ms. Dua berkas '
     'terakhir inilah buktinya: keduanya masuk ke cache peramban meskipun menunya tidak pernah '
     'ditekan. Sebaliknya, JadwalSeminarView.js dan PengaturanView.js sama sekali tidak diminta, '
     'yang menunjukkan bahwa prefetching bekerja secara terarah pada rute yang diprediksi paling '
     'sering dituju, bukan mengunduh seluruh halaman tanpa pandang bulu. Rekaman ini dihasilkan '
     'oleh skrip ukur_prefetch.cjs dan datanya tersimpan pada '
     'data_pengukuran/prefetch_network_log.json.'),
]

RUJUKAN = (
    ' Bukti bahwa mekanisme ini benar-benar berjalan pada aplikasi yang diuji disajikan pada '
    'Sub-bab 3.9.'
)


def kuning(par):
    for r in par.runs:
        if r.text.strip():
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def sisip(anchor, gaya_src, teks=None, align=None, gambar=None):
    """Sisipkan paragraf baru tepat setelah `anchor`, meniru gaya `gaya_src`."""
    el = copy.deepcopy(gaya_src._p)
    anchor._p.addnext(el)
    from docx.text.paragraph import Paragraph
    par = Paragraph(el, anchor._parent)
    for r in list(par.runs):
        r._element.getparent().remove(r._element)
    if gambar:
        par.add_run().add_picture(gambar, width=LEBAR)
    if teks:
        par.add_run(teks)
        kuning(par)
    if align is not None:
        par.paragraph_format.alignment = align
    return par


def main():
    doc = Document(DOCX)
    P = doc.paragraphs

    if any(p.text.strip().startswith(JUDUL) for p in P):
        print('Sub-bab 3.9 sudah ada. Tidak disisipkan ulang.')
        return

    gaya_judul = next(p for p in P if p.text.strip().startswith('3.8 Perbandingan Dampak'))
    gaya_isi = next(p for p in P if p.text.startswith('Kesimpulan dari tabel ini'))
    gaya_caption = next(p for p in P if p.text.strip().startswith('Gambar 3.4 Perbandingan'))
    gaya_gambar = P[P.index(gaya_caption) - 1]

    # titik sisip: paragraf terakhir Sub-bab 3.8
    anchor = gaya_isi

    anchor = sisip(anchor, gaya_judul, JUDUL)
    anchor.style = gaya_judul.style
    print('[OK] heading 3.9 disisipkan')

    anchor = sisip(anchor, gaya_isi, INTRO, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    for berkas, caption, uraian in GBR:
        path = os.path.join(IMG, berkas)
        if not os.path.exists(path):
            raise SystemExit(f'Gambar tidak ditemukan: {path}')
        anchor = sisip(anchor, gaya_gambar, gambar=path, align=WD_ALIGN_PARAGRAPH.CENTER)
        anchor = sisip(anchor, gaya_caption, caption, align=WD_ALIGN_PARAGRAPH.CENTER)
        anchor = sisip(anchor, gaya_isi, uraian, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        print(f'[OK] {berkas} + caption + uraian disisipkan')

    # rujukan maju dari Sub-bab 3.2
    for p in doc.paragraphs:
        if p.text.startswith('Dengan cara ini, ketika pengguna benar-benar menekan menu'):
            if 'Sub-bab 3.9' not in p.text:
                run = p.add_run(RUJUKAN)
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                print('[OK] rujukan ke Sub-bab 3.9 ditambahkan pada Sub-bab 3.2')
            break

    doc.save(DOCX)
    print('\nTersimpan.')


if __name__ == '__main__':
    main()
