# -*- coding: utf-8 -*-
"""Menyelaraskan penyebutan Supabase pada naskah dengan kondisi kode yang sebenarnya.

Lapisan akses data SIMTA diimplementasikan sebagai service layer asinkron yang
meniru pola Supabase (lihat src/services/supabase.js), bukan koneksi Supabase
daring: package.json tidak memuat @supabase/supabase-js dan bundle hasil build
tidak mengandung kode Supabase. Skrip ini memperbaiki kalimat-kalimat terkait
agar naskah konsisten dengan repositori.

Jalankan saat Revisi-Tesis-Mustari.docx tertutup:
    python perbaiki_supabase_docx.py
"""
import os
import shutil
from docx import Document

BASE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(BASE, 'Revisi-Tesis-Mustari.docx')
BACKUP = os.path.join(BASE, 'Revisi-Tesis-Mustari.SEBELUM-PERBAIKAN-SUPABASE.docx')

GANTI = [
    ('manajemen data global (Pinia), koneksi database (Supabase), dan otentikasi pengguna',
     'manajemen data global (Pinia), lapisan layanan data asinkron, dan otentikasi pengguna'),
    ('Pinia (state management), Supabase (koneksi database real-time), dan Vue Router dengan 15+ rute',
     'Pinia (state management), lapisan layanan data asinkron bergaya Supabase, dan Vue Router dengan 15+ rute'),
    ('Sistem Informasi Manajemen Tugas Akhir yang menggunakan Vue.js 3, Chart.js, Supabase, dan Pinia secara bersamaan',
     'Sistem Informasi Manajemen Tugas Akhir yang menggunakan Vue.js 3, Chart.js, Pinia, dan Vue Router secara '
     'bersamaan, dengan lapisan layanan data asinkron bergaya Supabase'),
    ('Vue-Router 4, Pinia Store v2, Chart.js, Supabase — pustaka pendukung SIMTA',
     'Vue-Router 4, Pinia Store v2, Chart.js, Tailwind CSS — pustaka pendukung SIMTA'),
    ('Ketergantungan pada 4 pustaka besar secara bersamaan: Chart.js, Pinia, Supabase, dan Vue-Router',
     'Ketergantungan pada 3 pustaka besar secara bersamaan: Chart.js, Pinia, dan Vue-Router'),
    ('Operasi CRUD dengan API asinkron',
     'Operasi CRUD melalui lapisan layanan data asinkron'),
    ('Autentikasi pengguna berbasis Supabase',
     'Alur autentikasi pengguna'),
    ('dan Supabase API sebagai lapisan data. Setiap perubahan data dari Supabase secara otomatis memperbarui',
     'dan lapisan layanan data asinkron (service layer) yang meniru pola akses Supabase. Setiap perubahan data '
     'dari lapisan layanan tersebut secara otomatis memperbarui'),
    ('komponen Vue, views, router, store Pinia, konfigurasi Supabase',
     'komponen Vue, views, router, store Pinia, dan lapisan layanan data asinkron'),
]

CATATAN_SETELAH = 'Kerumitan ini menjadi alasan utama mengapa pendekatan Lazy Loading diperlukan'
CATATAN = (
    'Perlu dicatat bahwa lapisan akses data pada SIMTA diimplementasikan sebagai service layer asinkron '
    'yang meniru pola pemanggilan Supabase (promise-based dengan penundaan jaringan tersimulasi), bukan '
    'koneksi ke basis data daring. Pilihan ini diambil agar pengujian sepenuhnya CPU-bound dan hasilnya '
    'dapat direproduksi tanpa bergantung pada ketersediaan layanan pihak ketiga maupun latensi jaringan. '
    'Karena strategi code splitting yang diteliti bekerja pada lapisan pemuatan modul JavaScript, keputusan '
    'ini tidak memengaruhi validitas perbandingan antara versi baseline dan optimized.'
)


def ganti_di_paragraf(par, cari, ganti):
    """Mengganti teks yang mungkin terpecah ke beberapa run, dengan tetap
    mempertahankan format run-run yang tidak tersentuh."""
    runs = par.runs
    if not runs:
        return False
    penuh = ''.join(r.text for r in runs)
    idx = penuh.find(cari)
    if idx == -1:
        return False

    # petakan posisi awal tiap run
    batas, pos = [], 0
    for r in runs:
        batas.append((pos, pos + len(r.text), r))
        pos += len(r.text)

    akhir = idx + len(cari)
    sudah_ditulis = False
    for mulai_r, akhir_r, r in batas:
        if akhir_r <= idx or mulai_r >= akhir:
            continue  # run di luar rentang, biarkan apa adanya
        sisa_kiri = r.text[:max(0, idx - mulai_r)]
        sisa_kanan = r.text[max(0, akhir - mulai_r):] if akhir_r > akhir else ''
        if not sudah_ditulis:
            r.text = sisa_kiri + ganti + sisa_kanan
            sudah_ditulis = True
        else:
            r.text = sisa_kiri + sisa_kanan
    return True


def semua_paragraf(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for baris in t.rows:
            for sel in baris.cells:
                for p in sel.paragraphs:
                    yield p


def main():
    if not os.path.exists(BACKUP):
        shutil.copy2(DOCX, BACKUP)
        print('Cadangan dibuat:', os.path.basename(BACKUP))

    doc = Document(DOCX)
    hitung = {c: 0 for c, _ in GANTI}
    for par in semua_paragraf(doc):
        for cari, ganti in GANTI:
            while ganti_di_paragraf(par, cari, ganti):
                hitung[cari] += 1

    # sisipkan catatan metodologis setelah paragraf penutup Sub-bab 2.6
    sudah_ada = any(p.text.strip().startswith('Perlu dicatat bahwa lapisan akses data')
                    for p in doc.paragraphs)
    if not sudah_ada:
        for p in doc.paragraphs:
            if p.text.strip().startswith(CATATAN_SETELAH):
                baru = doc.add_paragraph(CATATAN, style=p.style)
                p._p.addnext(baru._p)
                print('Catatan metodologis disisipkan pada Sub-bab 2.6.')
                break
    else:
        print('Catatan metodologis sudah ada, dilewati.')

    doc.save(DOCX)
    print('\nHasil penggantian:')
    for cari, jml in hitung.items():
        status = 'OK  ' if jml else 'TIDAK KETEMU'
        print('  [%s] %dx  %s...' % (status, jml, cari[:60]))


if __name__ == '__main__':
    main()
