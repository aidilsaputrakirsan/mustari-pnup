# -*- coding: utf-8 -*-
"""Menambahkan Lampiran 2 (ringkasan perubahan naskah) ke Revisi-Tesis-Mustari.docx.

Lampiran ini menjadi rujukan cepat bagi penguji untuk membandingkan naskah versi
ujian dengan naskah versi perbaikan, melengkapi sorotan kuning di BAB III dan IV.

Jalankan saat dokumen tertutup dari Microsoft Word:
    python tambah_lampiran_perubahan.py
"""
import os
import sys
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

if sys.stdout.encoding and sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

BASE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(BASE, 'Revisi-Tesis-Mustari.docx')

JUDUL = 'Lampiran 2. Ringkasan Perubahan Naskah Setelah Perbaikan Metodologi'


def format_tabel(t, lebar=9345):
    tblPr = t._tbl.tblPr
    w = tblPr.find(qn('w:tblW'))
    if w is None:
        w = OxmlElement('w:tblW')
        tblPr.append(w)
    w.set(qn('w:w'), str(lebar))
    w.set(qn('w:type'), 'dxa')

    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    tblPr.append(jc)

    borders = OxmlElement('w:tblBorders')
    for sisi in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + sisi)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'auto')
        borders.append(el)
    tblPr.append(borders)


def arsir(sel, warna='D9D9D9'):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), warna)
    sel._tc.get_or_add_tcPr().append(shd)


def buat_tabel(doc, judul_kolom, baris, ukuran=8):
    t = doc.add_table(rows=1, cols=len(judul_kolom))
    format_tabel(t)
    for sel, teks in zip(t.rows[0].cells, judul_kolom):
        sel.text = ''
        run = sel.paragraphs[0].add_run(teks)
        run.bold = True
        run.font.size = Pt(ukuran)
        arsir(sel)
    for data in baris:
        sels = t.add_row().cells
        for sel, teks in zip(sels, data):
            sel.text = ''
            run = sel.paragraphs[0].add_run(teks)
            run.font.size = Pt(ukuran)
    return t


TEMUAN = [
    ('1',
     'Kompresi Brotli tidak pernah terbentuk. Plugin vite-plugin-compression dipanggil dua kali '
     '(gzip dan brotli) dan berbagi cache internal, sehingga pemanggilan kedua melewati seluruh berkas.',
     'Berkas .br tidak pernah ada, sehingga manfaat kompresi tidak ikut terukur.',
     'Diganti plugin kustom compressAssets yang menjalankan gzip dan brotli dalam satu jalur per berkas.'),
    ('2',
     'Server pengujian tidak menyajikan berkas terkompresi. http-server dijalankan tanpa flag -g -b, '
     'sehingga selalu mengirim berkas mentah meskipun varian terkompresi tersedia.',
     'Peramban tidak pernah menerima berkas terkompresi, sehingga metrik jaringan tidak mencerminkan optimasi.',
     'Flag -g -b ditambahkan pada seluruh skrip pengukuran.'),
    ('3',
     'Lighthouse memakai Chrome sistem, bukan Chrome bawaan Puppeteer, sehingga muncul interstitial '
     'yang menggagalkan audit.',
     'Audit Lighthouse gagal pada seluruh percobaan sehingga hasilnya tidak dapat direproduksi.',
     'Skrip memaksa CHROME_PATH mengikuti Chrome bawaan Puppeteer.'),
]

ANGKA = [
    ('Lighthouse Performance Score (SIMTA)', 'Lighthouse',
     '66,2 menjadi 64,0 (turun 3,3%)', '56,8 menjadi 75,0 (naik 32,0%)',
     'Berbalik arah: optimasi kini terbukti menaikkan skor'),
    ('FCP SIMTA', 'Lighthouse',
     '5093,0 menjadi 5434,8 ms (memburuk)', '4960,4 menjadi 3250,2 ms (membaik 34,5%)',
     'Berbalik arah'),
    ('LCP SIMTA', 'Lighthouse',
     '5198,4 menjadi 5909,8 ms (memburuk)', '5228,4 menjadi 3790,8 ms (membaik 27,5%)',
     'Berbalik arah'),
    ('TTI SIMTA', 'Lighthouse',
     '5273,4 menjadi 5909,8 ms (memburuk)', '5566,8 menjadi 3964,0 ms (membaik 28,8%)',
     'Berbalik arah'),
    ('TBT SIMTA', 'Lighthouse',
     '105,2 menjadi 61,6 ms (membaik 41,4%)', '454,8 menjadi 315,4 ms (membaik 30,6%)',
     'Arah tetap membaik, besaran berubah'),
    ('FCP SIMTA kondisi normal', 'PerformanceObserver',
     '1144,0 menjadi 881,6 ms (membaik 22,9%)', '1103,2 menjadi 842,4 ms (membaik 23,6%)',
     'Arah dan besaran tetap serupa'),
    ('TBT SIMTA kondisi normal', 'PerformanceObserver',
     '111,8 menjadi 137,2 ms (memburuk)', '139,2 menjadi 104,6 ms (membaik 24,9%)',
     'Berbalik arah'),
    ('TBT SIMTA CPU diperlambat', 'PerformanceObserver',
     '1023,0 menjadi 790,8 ms (membaik 22,7%)', '710,4 menjadi 579,0 ms (membaik 18,5%)',
     'Arah tetap membaik'),
    ('FCP Company Profile CPU diperlambat', 'PerformanceObserver',
     '373,6 menjadi 486,4 ms (memburuk 30,2%)', '508,0 menjadi 508,8 ms (selisih 0,8 ms)',
     'Klaim degradasi 30,2% tidak dapat direproduksi'),
    ('FCP Company Profile', 'Lighthouse',
     '1352,8 menjadi 1579,0 ms (memburuk 16,7%)', '1361,4 menjadi 1578,2 ms (memburuk 15,9%)',
     'Konsisten pada dua pengukuran terpisah'),
]

ISI = [
    ('Sub-bab 3.1, mekanisme Chart.js',
     'Dinyatakan bahwa Chart.js hanya diunduh ketika pengguna membuka halaman yang menampilkan grafik.',
     'Dikoreksi: halaman Dashboard sebagai rute pertama sudah menampilkan grafik, sehingga vendor-chart.js '
     'tetap terunduh sejak awal. Manfaat sebenarnya berasal dari pengunduhan paralel, cache vendor '
     'terpisah, dan chunk yang tidak terunduh pada navigasi yang tidak dimulai dari Dashboard.'),
    ('Sub-bab 3.1, cuplikan kode',
     'Menampilkan pemanggilan viteCompression dua kali serta nama chunk vendor-charts dan vendor-core.',
     'Diselaraskan dengan kode repositori, yaitu plugin compressAssets serta nama chunk vendor-chart '
     'dan vendor-vue.'),
    ('Sub-bab 3.2, kelemahan dan solusi',
     'Hanya membahas lazy loading.',
     'Ditambahkan pembahasan kelemahan lazy loading berupa jeda pada setiap perpindahan halaman, '
     'beserta prefetching berbasis requestIdleCallback sebagai solusinya, lengkap dengan cuplikan kode.'),
    ('Sub-bab 3.4 sampai 3.6, interpretasi',
     'Menjelaskan trade-off bahwa FCP, LCP, dan TTI memburuk sebagai konsekuensi pemuatan bertahap.',
     'Penjelasan trade-off dihapus karena seluruh metrik SIMTA membaik setelah kompresi berfungsi.'),
    ('Sub-bab 3.8, klaim degradasi Company Profile',
     'Menyatakan FCP Company Profile memburuk 30,2% sebagai temuan.',
     'Ditambahkan catatan metodologis bahwa angka tersebut tidak dapat direproduksi pada '
     'PerformanceObserver. Kesimpulan dialihkan ke bukti Lighthouse yang stabil dan konsisten.'),
    ('BAB IV, kesimpulan',
     'Empat poin kesimpulan berbasis angka lama.',
     'Angka diperbarui dan ditambahkan satu poin mengenai peran prefetching sebagai penutup kelemahan '
     'lazy loading.'),
    ('BAB IV, saran penelitian lanjutan',
     'Lima butir saran.',
     'Ditambahkan satu butir, yaitu pengukuran kuantitatif jeda navigasi antar-halaman dengan dan '
     'tanpa prefetching.'),
    ('Gambar 3.4 sampai 3.9',
     'Grafik dibangkitkan dari data pengukuran lama.',
     'Seluruh grafik dibangkitkan ulang dari data pengukuran terbaru.'),
]

BATASAN = [
    'Manfaat prefetching terhadap jeda navigasi antar-halaman belum diukur secara kuantitatif. '
    'Penambahan prefetching dilandasi mekanisme kerjanya yang telah dipastikan berjalan pada kode, '
    'sementara biaya yang ditimbulkannya tetap ikut terekam pada metrik TBT dan penggunaan memori.',
    'Pada Company Profile, metrik PerformanceObserver terbukti sensitif terhadap variasi kondisi mesin '
    'antar-sesi pengukuran karena skala waktu muatnya hanya ratusan milidetik. Klaim persentase untuk '
    'aplikasi ini karenanya disandarkan pada data Lighthouse yang simpangan bakunya jauh lebih kecil.',
    'Konfigurasi build Company Profile tidak menyertakan kompresi Brotli maupun Gzip, sehingga '
    'perbaikan pada aplikasi tersebut tidak dapat dikaitkan dengan kompresi.',
]


def sudah_ada(doc):
    return any(p.text.strip().startswith('Lampiran 2.') for p in doc.paragraphs)


def main():
    doc = Document(DOCX)
    if sudah_ada(doc):
        print('Lampiran 2 sudah ada. Tidak ditambahkan ulang.')
        return

    doc.add_page_break()
    doc.add_heading(JUDUL, level=2)

    doc.add_paragraph(
        'Lampiran ini merangkum perbedaan antara naskah versi ujian dan naskah versi perbaikan. '
        'Seluruh bagian yang berubah pada BAB III dan BAB IV ditandai dengan sorotan berwarna kuning '
        'agar mudah ditelusuri. Perubahan berakar pada tiga temuan teknis pada pipeline pengukuran, '
        'bukan pada perubahan rancangan penelitian.'
    )

    doc.add_heading('A. Tiga Temuan yang Memicu Pengukuran Ulang', level=3)
    buat_tabel(doc, ('No', 'Temuan', 'Dampak pada Hasil Sebelumnya', 'Perbaikan'), TEMUAN)

    doc.add_heading('B. Perbandingan Angka Utama Sebelum dan Sesudah Perbaikan', level=3)
    doc.add_paragraph(
        'Seluruh angka berasal dari lima repetisi pengukuran pada kondisi dan perangkat yang sama.'
    )
    buat_tabel(doc, ('Metrik', 'Instrumen', 'Naskah Versi Ujian',
                     'Naskah Versi Perbaikan', 'Makna Perubahan'), ANGKA)

    doc.add_heading('C. Perubahan Isi Naskah', level=3)
    buat_tabel(doc, ('Bagian', 'Naskah Versi Ujian', 'Naskah Versi Perbaikan'), ISI)

    doc.add_heading('D. Batasan yang Disampaikan Secara Terbuka', level=3)
    for teks in BATASAN:
        p = doc.add_paragraph(teks, style='List Paragraph')
        p.paragraph_format.space_after = Pt(6)

    doc.save(DOCX)
    print('[OK] Lampiran 2 ditambahkan: A temuan, B perbandingan angka, C perubahan isi, D batasan.')


if __name__ == '__main__':
    main()
